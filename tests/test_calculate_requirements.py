import unittest
from collections import defaultdict
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

import calculate_requirements as calc


class ArithmeticTests(unittest.TestCase):
    def test_request_is_clamped_and_rounded_only_at_end(self):
        self.assertEqual(calc.request_amount(Decimal('12.01'), Decimal('10')), 3)
        self.assertEqual(calc.request_amount(Decimal('9.99'), Decimal('10')), 0)
        self.assertEqual(calc.request_amount(Decimal('10'), Decimal('10')), 0)

    def test_selects_most_recent_valid_bom(self):
        def bom(name, start, end=None):
            return calc.BomVersion('1200000', name, '', start, end, '')
        versions = {'old': bom('old', '14040101'),
                    'new': bom('new', '14050101'),
                    'expired': bom('expired', '14050301', '14050401'),
                    # Even an end date in the future must be excluded.
                    'dated_future': bom('dated_future', '14050701', '14060101')}
        chosen, candidates = calc.choose_version(versions, '14050809')
        self.assertEqual(chosen.bom, 'new')
        self.assertEqual({v.bom for v in candidates}, {'old', 'new'})
        with self.assertRaises(ValueError):
            calc.choose_version(versions, '14030101')
        with self.assertRaises(ValueError):
            calc.choose_version({'dated_future': versions['dated_future']}, '14050809')

    def test_loader_discards_every_bom_row_with_a_dated_end(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'bom.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.append(['header'] * 38)

            def add_row(bom, component, end):
                row = [None] * 38
                row[2], row[4], row[5], row[6] = '1200000', bom, 'name', 1
                row[10], row[11], row[12], row[13] = component, 'part', 1, 'عدد'
                row[26], row[33], row[34] = 50, '1405/01/01', end
                sheet.append(row)

            add_row('open', '1100001', None)
            # Filtering is row-by-row, including a dated row in the same BOM.
            add_row('open', '1100002', '1406/01/01')
            add_row('dated', '1100003', '1405/01/02')
            book.save(path)
            with patch.object(calc, 'BOM_FILE', path):
                versions = calc.load_boms({'1200000'})
            self.assertEqual(set(versions['1200000']), {'open'})
            self.assertEqual([x.component for x in versions['1200000']['open'].lines], ['1100001'])

    def test_formula_cache_survives_a_preceding_styled_blank_cell(self):
        book = Workbook()
        sheet = book.active
        sheet['H4'].fill = PatternFill('solid', fgColor='FFFFFF')
        cache = defaultdict(dict)
        calc.set_formula(sheet, 'I4', '=SUM(2,3)', Decimal(5), cache)
        with TemporaryDirectory() as directory:
            path = Path(directory) / 'result.xlsx'
            calc.save_with_cached_values(book, path, cache)
            with_formulas = load_workbook(path, read_only=True, data_only=False)
            with_values = load_workbook(path, read_only=True, data_only=True)
            self.assertEqual(with_formulas.active['I4'].value, '=SUM(2,3)')
            self.assertEqual(with_values.active['I4'].value, 5)
            with_formulas.close()
            with_values.close()

    def test_alternative_percentages_and_full_alternatives(self):
        def line(component, group, priority, pct):
            return calc.BomLine('1200000', 'bom', component, '', Decimal(1), Decimal(1), 'عدد',
                                group, priority, Decimal(pct), 2, '50')
        version = calc.BomVersion('1200000', 'bom', '', '14050101', None, '')
        version.lines = [line('1100001', 'A', 1, 50), line('1100002', 'A', 2, 50),
                         line('1100003', 'B', 2, 100), line('1100004', 'B', 1, 100)]
        shares = {line.component: share for line, share, _ in calc.split_bom_lines(version)}
        self.assertEqual(shares, {'1100001': Decimal('.5'), '1100002': Decimal('.5'),
                                  '1100003': Decimal(0), '1100004': Decimal(1)})


class SourceAndOutputTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = calc.build_data()
        cls.workbook, cls.cache = calc.create_workbook(cls.data)
        cls.temp_directory = TemporaryDirectory()
        cls.output = Path(cls.temp_directory.name) / 'request.xlsx'
        calc.save_with_cached_values(cls.workbook, cls.output, cls.cache)

    @classmethod
    def tearDownClass(cls):
        cls.temp_directory.cleanup()

    def test_plans_are_aggregated_and_versions_are_auditable(self):
        data = self.data
        self.assertEqual((data['start'], data['end']), ('14050809', '14050810'))
        self.assertEqual(len(data['products']), 13)
        self.assertEqual(data['products']['1200932']['old'], 650)
        self.assertEqual(data['products']['1200932']['new'], 800)  # two next-plan rows
        self.assertEqual(data['products']['1201643']['selected'], {'13416020183'})
        self.assertEqual(len(data['products']['1201643']['candidates']), 3)
        self.assertEqual(len(data['totals']), 313)
        self.assertEqual(len(data['detail']), 604)
        self.assertTrue(all(version.end is None for versions in data['versions'].values() for version in versions.values()))
        self.assertTrue(all(item['version'].end is None for item in data['detail']))

    def test_each_alternative_group_has_one_total_share(self):
        grouped = defaultdict(list)
        for item in self.data['detail']:
            line = item['line']
            if line.group:
                grouped[(line.product, line.bom, line.group)].append(item)
        self.assertTrue(grouped)
        for key, items in grouped.items():
            with self.subTest(key=key):
                self.assertEqual(sum((x['share'] for x in items), Decimal(0)), Decimal(1))
                # The what-if allocation must cover the scheduled product units
                # without counting both full alternatives. Division can leave a
                # tiny decimal-context residue.
                for period in ('old', 'new'):
                    actual = sum((x['priority_' + period] / x['line'].factor for x in items), Decimal(0))
                    self.assertLess(abs(actual - items[0]['product_' + period]), Decimal('1e-18'))

    def test_manual_component_cross_checks(self):
        data = self.data
        part = data['totals']['1100458']
        self.assertEqual((part['old'], part['new'], data['machine']['1100458'].quantity),
                         (Decimal(3000), Decimal(3000), Decimal(2992)))
        self.assertEqual(calc.request_amount(part['old'] + part['new'], data['machine']['1100458'].quantity), 3008)
        self.assertFalse(data['machine']['1110960'].saw_numeric)
        self.assertEqual(data['machine']['1110960'].quantity, 0)
        # The 12-prefix subassembly is not recursively exploded into another
        # unscheduled production plan or silently offset by planned completions.
        part = data['totals']['1201605']
        self.assertEqual((part['old'], part['new']), (Decimal(1460), Decimal(1200)))
        self.assertEqual(calc.request_amount(part['old'] + part['new'], data['machine']['1201605'].quantity), 2660)

    def test_workbook_formulas_and_unknown_fields(self):
        sheet = self.workbook['درخواست قطعات']
        self.assertEqual(sheet.max_row, 316)
        codes = set()
        for index, row in enumerate(sheet.iter_rows(min_row=4, values_only=True), 4):
            code, _, request, open_requests, _, _, stock, raw_stock, old, new, balance, next_only, priority, *_ = row
            self.assertNotIn(code, codes)
            codes.add(code)
            self.assertIsNone(open_requests)
            self.assertIsNone(raw_stock)
            self.assertEqual(request, f'=ROUNDUP(MAX(0,ROUND(I{index}+J{index}-G{index},10)),0)')
            self.assertTrue(old.startswith("=SUMIF('جزئیات BOM'!"))
            self.assertTrue(new.startswith("=SUMIF('جزئیات BOM'!"))
            self.assertEqual(balance, f'=ROUND(G{index}-I{index},10)')
            self.assertEqual(next_only, f'=ROUNDUP(MAX(0,ROUND(J{index}-G{index},10)),0)')
            self.assertIsInstance(priority, (int, float))  # explicitly labelled static scenario
        self.assertEqual(len(codes), 313)
        self.assertIn('ثابت', sheet['M3'].value)
        detail = self.workbook['جزئیات BOM']
        self.assertEqual(detail.max_row, 607)
        self.assertEqual(detail['E3'].value, 'تاریخ پایان اعتبار')
        for index, row in enumerate(detail.iter_rows(min_row=4, values_only=True), 4):
            self.assertIsNone(row[4])
            self.assertTrue(row[16].startswith("=SUMIFS('ورودی برنامه'!"))
            self.assertTrue(row[17].startswith("=SUMIFS('ورودی برنامه'!"))
            self.assertEqual(row[18], f'=ROUND(Q{index}*J{index}/K{index}*O{index}/100,10)')
            self.assertEqual(row[19], f'=ROUND(R{index}*J{index}/K{index}*O{index}/100,10)')
            if row[15] == 'سهم درصد استفاده BOM':
                self.assertEqual(row[14], f'=N{index}')

    def test_saved_workbook_has_formulas_and_correct_cached_values(self):
        formulas = load_workbook(self.output, read_only=True, data_only=False)
        values = load_workbook(self.output, read_only=True, data_only=True)
        self.assertEqual(values.sheetnames, ['درخواست قطعات', 'ورودی برنامه', 'برنامه محصولات ۱۲', 'جزئیات BOM', 'راهنما و فرض‌ها'])
        self.assertEqual(formulas.calculation.calcMode, 'auto')
        self.assertTrue(formulas.calculation.fullCalcOnLoad)
        self.assertTrue(formulas.calculation.forceFullCalc)
        self.assertIsInstance(values['درخواست قطعات']['C4'].value, int)
        for f_sheet, v_sheet in zip(formulas, values):
            for f_row, v_row in zip(f_sheet.iter_rows(), v_sheet.iter_rows()):
                for f_cell, v_cell in zip(f_row, v_row):
                    self.assertNotEqual(v_cell.data_type, 'e')
                    if f_cell.data_type == 'f':
                        self.assertIsInstance(v_cell.value, (int, float), f'{f_sheet.title}!{f_cell.coordinate}')
        for row in values['درخواست قطعات'].iter_rows(min_row=4, values_only=True):
            code, _, request, _, _, _, stock, _, old, new, balance, next_only, *_ = row
            part = self.data['totals'][str(code)]
            self.assertEqual(old, calc.number(part['old']))
            self.assertEqual(new, calc.number(part['new']))
            self.assertEqual(request, calc.request_amount(part['old'] + part['new'], Decimal(str(stock))))
            self.assertEqual(next_only, calc.request_amount(part['new'], Decimal(str(stock))))
        formulas.close()
        values.close()


if __name__ == '__main__':
    unittest.main()
