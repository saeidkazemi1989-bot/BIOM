#!/usr/bin/env python3
"""Create a traceable, provisional material-request workbook from the BIOM exports.

The source layouts and assumptions are deliberately explicit. In particular, this
is a planning estimate, not evidence that the earlier plan remains unconsumed or
that a particular BOM/alternative has been released by engineering.
"""
from __future__ import annotations

import math
import os
import re
import tempfile
from collections import defaultdict
from contextlib import closing
from dataclasses import dataclass, field
from decimal import Decimal, ROUND_CEILING
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
BOM_FILE = ROOT / "bom(1).xlsx"
PREV_FILE = ROOT / "برنامه 3روز.xlsx"
NEXT_FILE = ROOT / "برنامه بعدی.xlsx"
STOCK_FILE = ROOT / "موجودی به تفکیک انبار (1).xlsx"
OUTPUT_FILE = ROOT / "درخواست قطعات محاسبه شده.xlsx"
WAREHOUSE = "انبار پای کار مونتاژ ماشینی - جاده مخصوص"
ZERO = Decimal(0)
ONE = Decimal(1)
TRANSLATE_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def clean(value):
    return str(value if value is not None else "").translate(TRANSLATE_DIGITS).replace("\u200e", "").replace("\u200f", "").strip()


def code_of(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return clean(value)


def decimal_of(value, *, blank=None):
    if value is None or value == "":
        return blank
    if isinstance(value, bool):
        raise ValueError(f"Boolean is not a quantity: {value}")
    try:
        number = Decimal(clean(value).replace(",", ""))
    except Exception as exc:
        raise ValueError(f"Invalid numeric quantity: {value!r}") from exc
    if not number.is_finite():
        raise ValueError(f"Non-finite quantity: {value!r}")
    return number


def day_of(value):
    digits = re.sub(r"\D", "", clean(value))
    if len(digits) != 8:
        raise ValueError(f"Invalid Jalali date: {value!r}")
    return digits


def date_range_of(value):
    # Source contains e.g. 14050809-10, meaning 9th and 10th of that month.
    text = clean(value)
    match = re.fullmatch(r"(\d{8})\s*[-–]\s*(\d{2}|\d{8})", text)
    if not match:
        raise ValueError(f"Cannot read next-plan date range: {value!r}")
    start, end = match.groups()
    end = start[:6] + end if len(end) == 2 else end
    if end < start:
        raise ValueError(f"Invalid next-plan date range: {value!r}")
    return start, end


def fmt_day(value):
    return f"{value[:4]}/{value[4:6]}/{value[6:]}"


def number(value):
    value = Decimal(value)
    return int(value) if value == value.to_integral_value() else float(value)


def excel_code(value):
    return int(value) if value.isdecimal() else value


def request_amount(demand, stock):
    shortfall = demand - stock
    return int(shortfall.to_integral_value(rounding=ROUND_CEILING)) if shortfall > 0 else 0


@dataclass(frozen=True)
class PlanLine:
    code: str
    title: str
    quantity: Decimal
    day: str


@dataclass(frozen=True)
class BomLine:
    product: str
    bom: str
    component: str
    title: str
    coefficient: Decimal
    base: Decimal
    unit: str
    group: str
    priority: int | None
    pct: Decimal | None
    source_row: int
    warehouse: str

    @property
    def factor(self):
        return self.coefficient / self.base


@dataclass
class BomVersion:
    product: str
    bom: str
    title: str
    start: str
    end: str | None
    default_id: str
    lines: list[BomLine] = field(default_factory=list)

    def valid_on(self, day):
        return self.start <= day and (self.end is None or day <= self.end)


@dataclass
class Stock:
    quantity: Decimal = ZERO
    saw_numeric: bool = False
    saw_row: bool = False


def load_plans():
    old = []
    with closing(load_workbook(PREV_FILE, read_only=True, data_only=True)) as book:
        for row in book.active.iter_rows(min_row=2, values_only=True):
            product = code_of(row[1])
            if not product.startswith("12"):
                continue
            quantity = decimal_of(row[4])
            if quantity is None or quantity < ZERO:
                raise ValueError(f"Missing/negative prior plan quantity for {product}")
            old.append(PlanLine(product, clean(row[2]), quantity, day_of(row[0])))

    new = []
    with closing(load_workbook(NEXT_FILE, read_only=True, data_only=True)) as book:
        rows = book.active.iter_rows(values_only=True)
        start, end = date_range_of(next(rows)[2])
        next(rows)  # title row
        next(rows)  # headings
        for row in rows:
            product = code_of(row[11])
            if not product.startswith("12"):
                continue
            quantity = decimal_of(row[8])
            if quantity is None or quantity < ZERO:
                raise ValueError(f"Missing/negative next plan quantity for {product}")
            new.append(PlanLine(product, clean(row[10]), quantity, start))
    if not old or not new:
        raise ValueError("One or both plans contain no product codes starting with 12")
    return old, new, start, end


def load_boms(products):
    result = defaultdict(dict)
    with closing(load_workbook(BOM_FILE, read_only=True, data_only=True)) as book:
        for row_number, row in enumerate(book.active.iter_rows(min_row=2, values_only=True), 2):
            product = code_of(row[2])
            if product not in products:
                continue
            # Only rows whose "تاریخ پایان اعتبار" is empty are eligible.
            # Discard all dated end values, even when they are in the future.
            if clean(row[34]):
                continue
            bom = clean(row[4])
            start = day_of(row[33])
            end = None
            default_id = clean(row[36])
            versions = result[product]
            version = versions.setdefault(bom, BomVersion(product, bom, clean(row[5]), start, end, default_id))
            if (version.start, version.end, version.default_id) != (start, end, default_id):
                raise ValueError(f"Inconsistent dates/default ID within BOM {product}/{bom}")
            component = code_of(row[10])
            coefficient, base = decimal_of(row[12]), decimal_of(row[6])
            if not component or coefficient is None or coefficient < 0 or base is None or base <= 0:
                raise ValueError(f"Invalid BOM coefficient/base/component at row {row_number}")
            priority = int(row[18]) if row[18] not in (None, "") else None
            version.lines.append(BomLine(
                product, bom, component, clean(row[11]), coefficient, base,
                clean(row[13]), clean(row[17]), priority, decimal_of(row[20]),
                row_number, code_of(row[26]),
            ))
    return result


def choose_version(versions, day):
    candidates = [v for v in versions.values() if v.end is None and v.valid_on(day)]
    if not candidates:
        raise ValueError(f"No BOM with a blank end date valid on {fmt_day(day)}")
    # The export has a default-BOM ID but no ID for each version with which to
    # match it. The most recent valid start is therefore a *provisional* choice.
    candidates.sort(key=lambda v: (v.start, v.bom))
    return candidates[-1], candidates


def split_bom_lines(version):
    """Yield (line, applied share, explanation) without duplicating alternatives."""
    groups = defaultdict(list)
    for line in version.lines:
        if line.group:
            groups[line.group].append(line)
        else:
            yield line, ONE, "مصرف مستقیم"
    for group, lines in sorted(groups.items()):
        lines.sort(key=lambda x: (x.priority if x.priority is not None else math.inf, x.component, x.source_row))
        if any(line.priority is None or line.pct is None or line.pct < 0 for line in lines):
            raise ValueError(f"Missing substitution priority/percentage in {version.bom}, group {group}")
        total = sum((line.pct for line in lines), ZERO)
        if total == 100:
            for line in lines:
                yield line, line.pct / 100, "سهم درصد استفاده BOM"
        elif len(lines) > 1 and all(line.pct == 100 for line in lines) and lines[0].priority < lines[1].priority:
            # 100%+100% is not 200% consumption: they are full alternatives.
            for i, line in enumerate(lines):
                yield line, ONE if i == 0 else ZERO, "۱۰۰٪ تکراری؛ فقط اولویت اول"
        else:
            raise ValueError(f"Unsupported substitution percentages in {version.bom}, group {group}: {total}")


def load_stock():
    machine = defaultdict(Stock)
    other = defaultdict(list)
    with closing(load_workbook(STOCK_FILE, read_only=True, data_only=True)) as book:
        for raw_code, title, warehouse, raw_quantity in book.active.iter_rows(min_row=2, values_only=True):
            code = code_of(raw_code)
            if not code:
                continue
            warehouse = clean(warehouse)
            quantity = decimal_of(raw_quantity)
            if warehouse == WAREHOUSE:
                entry = machine[code]
                entry.saw_row = True
                if quantity is not None:
                    entry.saw_numeric = True
                    entry.quantity += quantity
            elif quantity not in (None, ZERO):
                other[code].append((warehouse, quantity))
    return machine, other


def build_data():
    old, new, start, end = load_plans()
    products = {p.code for p in old + new}
    versions = load_boms(products)
    machine, other = load_stock()
    quantities = defaultdict(lambda: {"old": ZERO, "new": ZERO})
    product_info = {p.code: {"title": p.title, "old": ZERO, "new": ZERO,
                              "selected": set(), "candidates": set(), "default_ids": set()}
                    for p in old + new}
    for period, lines in (("old", old), ("new", new)):
        for plan in lines:
            if plan.code not in versions:
                raise ValueError(f"No BOM found for {plan.code}")
            version, candidates = choose_version(versions[plan.code], plan.day)
            if period == "new":
                # The next plan has one total for two days; do not silently pick
                # one BOM if the applicable version differs between those days.
                end_version, end_candidates = choose_version(versions[plan.code], end)
                if version.bom != end_version.bom:
                    raise ValueError(f"Next-plan BOM changes within its date range for {plan.code}")
                candidates = list({v.bom: v for v in candidates + end_candidates}.values())
            quantities[(plan.code, version.bom)][period] += plan.quantity
            p = product_info[plan.code]
            p[period] += plan.quantity
            p["selected"].add(version.bom)
            p["candidates"].update(v.bom for v in candidates)
            p["default_ids"].add(version.default_id)

    detail = []
    totals = defaultdict(lambda: {"title": "", "unit": set(), "old": ZERO, "new": ZERO,
                                  "priority_old": ZERO, "priority_new": ZERO,
                                  "priorities": set(), "substitution": False,
                                  "full_alternatives": False})
    for key in sorted(quantities):
        product, bom = key
        version = versions[product][bom]
        qty = quantities[key]
        for line, share, method in split_bom_lines(version):
            if line.warehouse != "50":
                raise ValueError(f"Unexpected BOM work-warehouse {line.warehouse} at row {line.source_row}")
            if "عدد" not in line.unit:
                raise ValueError(f"Non-counting unit requires separate rounding at row {line.source_row}: {line.unit}")
            item = {"line": line, "version": version, "share": share, "method": method,
                    "product_old": qty["old"], "product_new": qty["new"],
                    "old": qty["old"] * line.factor * share,
                    "new": qty["new"] * line.factor * share,
                    "priority_old": ZERO, "priority_new": ZERO}
            detail.append(item)
            total = totals[line.component]
            total["title"] = total["title"] or line.title
            total["unit"].add(line.unit)
            total["old"] += item["old"]
            total["new"] += item["new"]
            if line.group:
                total["substitution"] = True
                total["priorities"].add(line.priority)
                total["full_alternatives"] |= "۱۰۰٪" in method

    # Optional, separate what-if: use available substitutes by priority. Fixed
    # requirements get stock first, then old-plan groups, then next-plan groups.
    # This is greedy, not an optimal allocation, and ignores BOM use percentages.
    remaining = {component: max(ZERO, machine[component].quantity) for component in totals}
    grouped = defaultdict(list)
    for item in detail:
        line = item["line"]
        if not line.group:
            for period in ("old", "new"):
                item["priority_" + period] = item[period]
            remaining[line.component] = max(ZERO, remaining[line.component] - item["old"] - item["new"])
        else:
            grouped[(line.product, line.bom, line.group)].append(item)
    for period in ("old", "new"):
        for key in sorted(grouped):
            candidates = sorted(grouped[key], key=lambda x: (x["line"].priority, x["line"].component, x["line"].source_row))
            unfilled = candidates[0]["product_" + period]
            for item in candidates:
                line = item["line"]
                if line.factor <= 0:
                    continue
                units = min(unfilled, remaining[line.component] / line.factor)
                part_quantity = units * line.factor
                item["priority_" + period] += part_quantity
                remaining[line.component] -= part_quantity
                unfilled -= units
            if unfilled > 0:
                first = candidates[0]
                first["priority_" + period] += unfilled * first["line"].factor
    for item in detail:
        total = totals[item["line"].component]
        total["priority_old"] += item["priority_old"]
        total["priority_new"] += item["priority_new"]

    return {"old": old, "new": new, "start": start, "end": end,
            "products": product_info, "versions": versions, "quantities": quantities,
            "machine": machine, "other": other, "totals": totals, "detail": detail}


NAVY = "183153"
BLUE = "24527A"
TEAL = "DDEFF0"
PALE_BLUE = "EAF2FA"
PALE_YELLOW = "FFF2CC"
PALE_RED = "FCE8E6"
GREEN = "E4F4EB"
WHITE = "FFFFFF"
GREY = "F3F5F7"
BORDER = Border(bottom=Side(style="hair", color="D7DEE6"))


def setup_table(sheet, title, subtitle, headers, widths):
    sheet.sheet_view.rightToLeft = True
    n = len(headers)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    sheet.cell(1, 1, title)
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].font = Font(name="Tahoma", size=13, bold=True, color=WHITE)
    sheet["A1"].alignment = Alignment(horizontal="right", vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    sheet.cell(2, 1, subtitle)
    sheet["A2"].fill = PatternFill("solid", fgColor=PALE_YELLOW)
    sheet["A2"].font = Font(name="Tahoma", size=10, bold=True, color="7A5100")
    sheet["A2"].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 30
    for col, (header, width) in enumerate(zip(headers, widths), 1):
        cell = sheet.cell(3, col, header)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Tahoma", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.row_dimensions[3].height = 40
    sheet.freeze_panes = "C4"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.print_title_rows = "1:3"


def add_row(sheet, values):
    sheet.append(values)
    idx = sheet.max_row
    for cell in sheet[idx]:
        cell.font = Font(name="Tahoma", size=9, color="263445")
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", horizontal="right" if isinstance(cell.value, str) else "center")
        if idx % 2 == 0:
            cell.fill = PatternFill("solid", fgColor="F7FAFC")
        if isinstance(cell.value, (int, float)):
            cell.number_format = '#,##0.##;[Red]-#,##0.##;0'
    sheet.row_dimensions[idx].height = 20


def auto_filter(sheet):
    sheet.auto_filter.ref = f"A3:{get_column_letter(sheet.max_column)}{sheet.max_row}"


def set_formula(sheet, coordinate, formula, cached_value, cache):
    """Keep a real Excel formula and the already-validated starting result."""
    if not formula.startswith("="):
        raise ValueError(f"Not a formula: {formula!r}")
    sheet[coordinate] = formula
    cache[sheet.title][coordinate] = number(cached_value)


def save_with_cached_values(workbook, output_path, cache):
    """Save formulas plus numeric caches so viewers show values before recalculation.

    openpyxl writes formulas but not their cached values. The caches supplied here
    are computed from the same source records as the formulas. Excel recalculates
    the formulas on open or on input changes; this does not replace any formula.
    """
    output_path = Path(output_path)
    unknown = set(cache) - set(workbook.sheetnames)
    if unknown:
        raise ValueError(f"Unknown formula-cache sheets: {unknown}")
    # Styled blank cells can be self-closing; never consume the next formula.
    cell_pattern = re.compile(rb'(<c\b[^>]*\br="([A-Z]+[0-9]+)"[^>]*(?<!/)>)(.*?)(</c>)', re.DOTALL)
    value_pattern = re.compile(rb'<v(?:\s*/>|\s*>.*?</v>)', re.DOTALL)
    with tempfile.TemporaryDirectory(prefix=".formula-save-", dir=output_path.parent) as directory:
        original = Path(directory) / "openpyxl.xlsx"
        final = Path(directory) / "cached.xlsx"
        workbook.save(original)
        with ZipFile(original, "r") as source, ZipFile(final, "w") as destination:
            for info in source.infolist():
                payload = source.read(info.filename)
                for index, sheet_name in enumerate(workbook.sheetnames, 1):
                    if info.filename != f"xl/worksheets/sheet{index}.xml" or sheet_name not in cache:
                        continue
                    sheet_cache = cache[sheet_name]
                    written = set()

                    def replace_cell(match):
                        coordinate = match.group(2).decode("ascii")
                        if coordinate not in sheet_cache:
                            return match.group(0)
                        body = match.group(3)
                        if not re.search(rb'<f(?:\s[^>]*)?>', body):
                            raise ValueError(f"Cache points to non-formula cell {sheet_name}!{coordinate}")
                        replacement = b"<v>" + str(sheet_cache[coordinate]).encode("ascii") + b"</v>"
                        body, count = value_pattern.subn(replacement, body, count=1)
                        if count != 1:
                            raise ValueError(f"Missing value node in {sheet_name}!{coordinate}")
                        written.add(coordinate)
                        return match.group(1) + body + match.group(4)

                    payload = cell_pattern.sub(replace_cell, payload)
                    if written != set(sheet_cache):
                        raise ValueError(f"Unwritten formula caches in {sheet_name}: {set(sheet_cache) - written}")
                    break
                destination.writestr(info, payload)
        os.replace(final, output_path)


def create_workbook(data):
    wb = Workbook()
    cache = defaultdict(dict)
    sheet = wb.active
    sheet.title = "درخواست قطعات"
    headers = ["کد کالای مصرفی", "عنوان کالای مصرفی", "مقدار درخواست", "درخواست باز", "جایگزینی", "اولویت",
               "موجودی انبار پای کار", "موجودی انبار مواد اولیه", "نیاز برنامه سه‌روز", "نیاز برنامه بعدی",
               "مانده پس از سه‌روز", "درخواست فقط برنامه بعدی", "درخواست اولویت و جایگزین (ثابت)",
               "وضعیت موجودی پای کار", "واحد", "توضیحات", "برنامه تولید خود قطعه ۱۲ - قبل", "برنامه تولید خود قطعه ۱۲ - بعد"]
    widths = [17, 46, 18, 15, 13, 14, 19, 20, 19, 19, 19, 21, 26, 25, 15, 43, 21, 21]
    machine, totals = data["machine"], data["totals"]
    plan_end = 3 + len(data["quantities"])
    detail_end = 3 + len(data["detail"])
    plan_codes = f"'ورودی برنامه'!$A$4:$A${plan_end}"
    plan_boms = f"'ورودی برنامه'!$C$4:$C${plan_end}"
    plan_old = f"'ورودی برنامه'!$D$4:$D${plan_end}"
    plan_new = f"'ورودی برنامه'!$E$4:$E${plan_end}"
    detail_codes = f"'جزئیات BOM'!$G$4:$G${detail_end}"
    detail_old = f"'جزئیات BOM'!$S$4:$S${detail_end}"
    detail_new = f"'جزئیات BOM'!$T$4:$T${detail_end}"
    setup_table(sheet, "برآورد درخواست قطعات | برنامه قبلی و برنامه بعدی",
                "فقط BOM با پایان اعتبار خالی؛ برنامه قبلی مصرف‌نشده فرض شده است. ستون C فرمول دارد و سناریوی M ثابت است.", headers, widths)
    sheet["C3"].comment = Comment("فرمول: درخواست = ROUNDUP(MAX(0, نیاز سه‌روز + نیاز بعدی − موجودی پای کار), 0). گردکردن خطای شناوری قبل از ROUNDUP انجام می‌شود.", "راهنما")
    sheet["I3"].comment = Comment("جمع نیاز این قطعه در همه محصولات از ستون S در صفحه جزئیات BOM با SUMIF.", "راهنما")
    sheet["J3"].comment = Comment("جمع نیاز این قطعه در همه محصولات از ستون T در صفحه جزئیات BOM با SUMIF.", "راهنما")
    sheet["M3"].comment = Comment("این سناریوی کمکی فقط مقدار اولیه ثابت است؛ با تغییر برنامه، ضریب یا موجودی به‌روز نمی‌شود. برای محاسبه مجدد باید فایل را دوباره تولید کرد.", "راهنما")
    sheet["M3"].fill = PatternFill("solid", fgColor=PALE_YELLOW)
    sheet["M3"].font = Font(name="Tahoma", size=10, bold=True, color="7A5100")
    sorted_codes = sorted(totals, key=lambda c: (-request_amount(totals[c]["old"] + totals[c]["new"], machine[c].quantity), c))
    for code in sorted_codes:
        t, st = totals[code], machine[code]
        requested = request_amount(t["old"] + t["new"], st.quantity)
        only_new = request_amount(t["new"], st.quantity)
        alternate = request_amount(t["priority_old"] + t["priority_new"], st.quantity)
        status = "ثبت عددی" if st.saw_numeric else ("خالی در فایل؛ فرض صفر" if st.saw_row else "بدون ردیف؛ فرض صفر")
        priorities = "، ".join(str(x) for x in sorted(t["priorities"])) if t["priorities"] else "-"
        notes = []
        if t["substitution"]:
            notes.append("سهم جایگزین طبق BOM؛ نیازمند تأیید")
        if t["full_alternatives"]:
            notes.append("۱۰۰٪ تکراری: اولویت اول")
        if code.startswith("12"):
            notes.append("زیرمجموعه تولیدی؛ تکمیل برنامه آن تأیید نشده")
        if not st.saw_numeric:
            notes.append("موجودی نامشخص؛ صفر فرض شده")
        plan = data["products"].get(code)
        add_row(sheet, [excel_code(code), t["title"], requested, None, "دارد" if t["substitution"] else "ندارد",
                        priorities, number(st.quantity), None, number(t["old"]), number(t["new"]),
                        number(st.quantity - t["old"]), only_new, alternate, status, "، ".join(sorted(t["unit"])),
                        " | ".join(notes), number(plan["old"]) if plan else None, number(plan["new"]) if plan else None])
        row = sheet.max_row
        set_formula(sheet, f"I{row}", f"=SUMIF({detail_codes},A{row},{detail_old})", t["old"], cache)
        set_formula(sheet, f"J{row}", f"=SUMIF({detail_codes},A{row},{detail_new})", t["new"], cache)
        # Remove floating-point noise before rounding the final shortage up.
        set_formula(sheet, f"C{row}", f"=ROUNDUP(MAX(0,ROUND(I{row}+J{row}-G{row},10)),0)", requested, cache)
        set_formula(sheet, f"K{row}", f"=ROUND(G{row}-I{row},10)", st.quantity - t["old"], cache)
        set_formula(sheet, f"L{row}", f"=ROUNDUP(MAX(0,ROUND(J{row}-G{row},10)),0)", only_new, cache)
        if plan:
            set_formula(sheet, f"Q{row}", f"=SUMIF({plan_codes},A{row},{plan_old})", plan["old"], cache)
            set_formula(sheet, f"R{row}", f"=SUMIF({plan_codes},A{row},{plan_new})", plan["new"], cache)
        sheet.cell(row, 3).font = Font(name="Tahoma", bold=True, size=10, color=BLUE)
        if not st.saw_numeric:
            sheet.cell(row, 14).fill = PatternFill("solid", fgColor=PALE_YELLOW)
        if code.startswith("12"):
            sheet.cell(row, 16).fill = PatternFill("solid", fgColor=PALE_YELLOW)
        for col in (4, 8):
            sheet.cell(row, col).fill = PatternFill("solid", fgColor=GREY)
    sheet.conditional_formatting.add(
        f"C4:C{sheet.max_row}",
        CellIsRule(operator="greaterThan", formula=["0"],
                   fill=PatternFill("solid", fgColor=PALE_RED),
                   font=Font(name="Tahoma", bold=True, color="9C1C1C")),
    )
    auto_filter(sheet)

    inputs = wb.create_sheet("ورودی برنامه")
    input_headers = ["کد محصول", "نام محصول", "شماره BOM انتخاب‌شده", "تعداد برنامه سه‌روز", "تعداد برنامه بعدی", "شروع اعتبار", "تاریخ پایان اعتبار"]
    setup_table(inputs, "ورودی برنامه به تفکیک محصول و BOM",
                "برای آزمایش یا اصلاح محاسبه، فقط تعداد برنامه‌ها را در ستون‌های آبی ویرایش کنید. نسخه BOM به‌طور خودکار عوض نمی‌شود.",
                input_headers, [17, 48, 25, 23, 23, 17, 20])
    for (code, bom), quantities in sorted(data["quantities"].items()):
        version = data["versions"][code][bom]
        add_row(inputs, [excel_code(code), data["products"][code]["title"], bom,
                         number(quantities["old"]), number(quantities["new"]),
                         fmt_day(version.start), fmt_day(version.end) if version.end else None])
        for col in (4, 5):
            inputs.cell(inputs.max_row, col).fill = PatternFill("solid", fgColor=PALE_BLUE)
    inputs["D3"].comment = Comment("تعداد باقیمانده از برنامه سه‌روز؛ با تغییر آن سناریوی اصلی با همان BOM باز حساب می‌شود.", "راهنما")
    inputs["E3"].comment = Comment("تعداد برنامه بعدی؛ برای محاسبه جدید آن را در این صفحه ویرایش کنید.", "راهنما")
    validation = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=False)
    validation.showErrorMessage = True
    validation.errorTitle = "مقدار نامعتبر"
    validation.error = "تعداد برنامه باید عدد صفر یا بیشتر باشد."
    inputs.add_data_validation(validation)
    validation.add(f"D4:E{inputs.max_row}")
    auto_filter(inputs)

    products = wb.create_sheet("برنامه محصولات ۱۲")
    p_headers = ["کد محصول", "نام محصول", "برنامه سه‌روز", "برنامه بعدی", "موجودی محصول در انبار پای کار ماشینی",
                 "موجودی محصول در سایر پای کارها (اطلاعاتی)", "جزئیات سایر انبارها", "نسخه BOM انتخاب‌شده",
                 "نسخه‌های همزمان با پایان اعتبار خالی", "شناسه پیش‌فرض ثبت‌شده", "هشدار"]
    setup_table(products, "کنترل محصولات برنامه‌ریزی‌شده با کد ۱۲",
                "برنامه‌ها با فرمول از صفحه ورودی برنامه جمع می‌شوند. موجودی محصول نهایی اطلاعاتی است و از تعداد برنامه کم نشده است.",
                p_headers, [17, 48, 17, 17, 22, 25, 65, 25, 40, 22, 45])
    for code, p in sorted(data["products"].items()):
        other_rows = data["other"].get(code, [])
        other_total = sum((v for _, v in other_rows), ZERO)
        flags = []
        if len(p["candidates"]) > 1:
            flags.append("چند نسخه با پایان خالی؛ انتخاب جدیدترین شروع")
        if not any(p["default_ids"]):
            flags.append("شناسه BOM پیش‌فرض خالی است")
        if not machine[code].saw_numeric:
            flags.append("موجودی ماشینی خالی؛ صفر فرض شده")
        add_row(products, [excel_code(code), p["title"], number(p["old"]), number(p["new"]),
                           number(machine[code].quantity), number(other_total),
                           " | ".join(f"{w}: {number(q):,}" for w, q in other_rows),
                           "، ".join(sorted(p["selected"])), "، ".join(sorted(p["candidates"])),
                           "، ".join(sorted(x for x in p["default_ids"] if x)) or "ثبت نشده", " | ".join(flags)])
        row = products.max_row
        set_formula(products, f"C{row}", f"=SUMIF({plan_codes},A{row},{plan_old})", p["old"], cache)
        set_formula(products, f"D{row}", f"=SUMIF({plan_codes},A{row},{plan_new})", p["new"], cache)
        if flags:
            products.cell(row, 11).fill = PatternFill("solid", fgColor=PALE_YELLOW)
    auto_filter(products)

    detail = wb.create_sheet("جزئیات BOM")
    d_headers = ["کد محصول", "نام محصول", "شماره BOM", "شروع اعتبار", "تاریخ پایان اعتبار", "ردیف فایل BOM", "کد قطعه",
                 "عنوان قطعه", "واحد", "ضریب", "مبنای BOM", "گروه جایگزین", "اولویت", "درصد استفاده BOM",
                 "سهم اعمال شده (٪)", "روش محاسبه", "برنامه سه‌روز محصول", "برنامه بعدی محصول", "نیاز سه‌روز قطعه",
                 "نیاز بعدی قطعه", "نیاز سه‌روز اولویت (ثابت)", "نیاز بعدی اولویت (ثابت)"]
    setup_table(detail, "ریز فرمول محاسبه به تفکیک محصول و ردیف BOM",
                "ستون‌های Q تا T فرمول دارند. ستون‌های U و V سناریوی جایگزین ثابت‌اند و بعد از تغییر ورودی به‌روز نمی‌شوند.",
                d_headers, [16, 45, 24, 16, 16, 15, 16, 48, 16, 14, 14, 15, 12, 16, 19, 29, 19, 19, 19, 19, 23, 23])
    detail["O3"].comment = Comment("برای قطعات عادی ۱۰۰٪ است. در گروه‌های درصدی از ستون N می‌آید؛ در گروه ۱۰۰٪/۱۰۰٪ سهم انتخاب‌شده ۱۰۰ یا صفر و ثابت است.", "راهنما")
    detail["S3"].comment = Comment("نیاز سه‌روز = برنامه سه‌روز محصول × ضریب ÷ مبنای BOM × سهم اعمال‌شده ÷ ۱۰۰.", "راهنما")
    detail["T3"].comment = Comment("نیاز برنامه بعدی = برنامه بعدی محصول × ضریب ÷ مبنای BOM × سهم اعمال‌شده ÷ ۱۰۰.", "راهنما")
    for col in (21, 22):
        detail.cell(3, col).fill = PatternFill("solid", fgColor=PALE_YELLOW)
        detail.cell(3, col).font = Font(name="Tahoma", size=10, bold=True, color="7A5100")
    for item in sorted(data["detail"], key=lambda x: (x["line"].product, x["line"].bom, x["line"].source_row)):
        line, version = item["line"], item["version"]
        add_row(detail, [excel_code(line.product), data["products"][line.product]["title"], line.bom,
                         fmt_day(version.start), fmt_day(version.end) if version.end else None, line.source_row,
                         excel_code(line.component), line.title, line.unit, number(line.coefficient), number(line.base),
                         line.group or "-", line.priority, number(line.pct) if line.pct is not None else None,
                         number(item["share"] * 100), item["method"], number(item["product_old"]), number(item["product_new"]),
                         number(item["old"]), number(item["new"]), number(item["priority_old"]), number(item["priority_new"])])
        row = detail.max_row
        if item["method"] == "سهم درصد استفاده BOM":
            set_formula(detail, f"O{row}", f"=N{row}", item["share"] * 100, cache)
        set_formula(detail, f"Q{row}", f"=SUMIFS({plan_old},{plan_codes},A{row},{plan_boms},C{row})", item["product_old"], cache)
        set_formula(detail, f"R{row}", f"=SUMIFS({plan_new},{plan_codes},A{row},{plan_boms},C{row})", item["product_new"], cache)
        set_formula(detail, f"S{row}", f"=ROUND(Q{row}*J{row}/K{row}*O{row}/100,10)", item["old"], cache)
        set_formula(detail, f"T{row}", f"=ROUND(R{row}*J{row}/K{row}*O{row}/100,10)", item["new"], cache)
    auto_filter(detail)

    guide = wb.create_sheet("راهنما و فرض‌ها")
    guide.sheet_view.rightToLeft = True
    guide.column_dimensions["A"].width = 27
    guide.column_dimensions["B"].width = 125
    guide.merge_cells("A1:B1")
    guide["A1"] = "راهنمای برآورد و مواردی که باید قبل از ثبت درخواست تأیید شوند"
    guide["A1"].fill = PatternFill("solid", fgColor=NAVY)
    guide["A1"].font = Font(name="Tahoma", size=13, color=WHITE, bold=True)
    guide["A1"].alignment = Alignment(horizontal="right", vertical="center")
    guide.row_dimensions[1].height = 32
    min_old = min(x.day for x in data["old"])
    max_old = max(x.day for x in data["old"])
    missing = sum(not machine[c].saw_numeric for c in totals)
    overlap = [c for c, p in data["products"].items() if len(p["candidates"]) > 1]
    lines = [
        ("دیدن فرمول", "در اکسل روی عدد دلخواه در ستون «مقدار درخواست» یا روی سلول‌های نیاز در صفحه «جزئیات BOM» کلیک کنید؛ فرمول واقعی در نوار فرمول بالای صفحه نمایش داده می‌شود."),
        ("ویرایش برنامه", "تعدادهای سه‌روز و بعدی در ستون‌های آبی صفحه «ورودی برنامه» قابل ویرایش‌اند. با تغییر آنها و محاسبه خودکار اکسل، نیاز هر محصول و جمع نیاز و درخواست اصلی به‌روز می‌شود. موجودی در ستون G صفحه درخواست نیز قابل ویرایش است."),
        ("محدوده به‌روزرسانی", "فرمول‌ها فقط ردیف‌ها و BOMهای انتخاب‌شده در همین فایل را محاسبه می‌کنند. افزودن محصول، تغییر دوره یا انتخاب نسخه جدید BOM به بازتولید فایل نیاز دارد. اکسل را در حالت محاسبه خودکار نگه دارید."),
        ("ماهیت خروجی", "پیش‌نویس برآورد است، نه تأیید اجرای تولید، موجودی انبار مواد یا درخواست باز. در صفحه اول ستون‌های نمونه خروجی حفظ شده‌اند."),
        ("محصولات", "فقط کد محصول‌هایی که با ۱۲ شروع می‌شوند از برنامه‌ها انتخاب شده‌اند. کدهای خروجی، قطعات مصرفی BOM هستند و لزوماً با ۱۲ شروع نمی‌شوند."),
        ("دوره برنامه‌ها", f"برنامه سه‌روز از {fmt_day(min_old)} تا {fmt_day(max_old)}؛ برنامه بعدی از {fmt_day(data['start'])} تا {fmt_day(data['end'])}. ردیف‌های تکراری یک محصول جمع شده‌اند."),
        ("فرمول اصلی", "در صفحه درخواست، برای نمونه فرمول C4 برابر =ROUNDUP(MAX(0,ROUND(I4+J4-G4,10)),0) است: جمع نیاز سه‌روز و بعدی منهای موجودی، حداقل صفر و گرد به بالا. ROUND داخلی فقط خطای نمایش اعشار رایانه را حذف می‌کند."),
        ("هشدار مهم برنامه قبلی", "چون مقدار واقعی تولید/مصرف‌شده در فایل‌ها نیست، کل برنامه سه‌روز مصرف‌نشده و نیازمند رزرو فرض شده است. اگر بخشی قبلاً مصرف شده، این برآورد ممکن است بیش‌درخواست بدهد. ستون «درخواست فقط برنامه بعدی» سناریوی دیگر است؛ مقادیر دو ستون را با هم جمع نکنید."),
        ("محاسبه نیاز", "در صفحه جزئیات BOM برای هر ردیف، نیاز سه‌روز = برنامه محصول × ضریب ÷ مبنای BOM × سهم اعمال‌شده ÷ ۱۰۰ است؛ مثلاً S4 =ROUND(Q4*J4/K4*O4/100,10). ستون‌های I و J صفحه درخواست، نیاز همه محصولات را با SUMIF جمع می‌کنند."),
        ("نسخه BOM", "فقط ردیف‌هایی لحاظ شده‌اند که ستون «تاریخ پایان اعتبار» آنها خالی است؛ ردیف‌های دارای تاریخ پایان، حتی تاریخ آینده، حذف شده‌اند. از نسخه‌های باقی‌مانده که شروعشان تا تاریخ برنامه است، جدیدترین تاریخ شروع انتخاب شده است. شناسه هر نسخه برای تطبیق با شناسه پیش‌فرض موجود نیست؛ انتخاب پیش‌فرض قطعی نیست."),
        ("نسخه‌های همزمان", "برای محصولات " + ("، ".join(sorted(overlap)) if overlap else "هیچ‌کدام") + " چند نسخه با تاریخ پایان خالی و شروع معتبر همزمان وجود دارد. برای جلوگیری از دوباره‌شماری، فقط جدیدترین نسخه انتخاب شده؛ انتخاب باید توسط مهندسی تأیید شود."),
        ("قطعات جایگزین", "در گروه‌هایی که درصد استفاده جمعاً ۱۰۰ است، سهم هر قطعه مطابق درصد BOM فرض شده است. اگر هر دو جایگزین ۱۰۰٪ دارند، فقط اولویت اول برای سناریوی اصلی منظور شده تا مصرف دوبرابر نشود. این فرض و مجاز بودن نسبت‌ها باید تأیید شود."),
        ("سناریوی دیگر جایگزین", "ستون M صفحه درخواست و ستون‌های U و V جزئیات BOM مقادیر اولیه ثابت هستند، نه فرمول. با تغییر برنامه، ضریب یا موجودی به‌روز نمی‌شوند و برای محاسبه مجدد آنها باید فایل را دوباره تولید کرد. این سناریوی تقریبی، موجودی جایگزین‌ها را بر اساس اولویت استفاده می‌کند و لزوماً درصدهای BOM را رعایت نمی‌کند."),
        ("پای کار", "فقط انبار پای کار مونتاژ ماشینی - جاده مخصوص محاسبه شده است؛ انبار پای کار در BOMهای انتخابی کد ۵۰ است. سایر انبارها برای درخواست قابل برداشت فرض نشده‌اند."),
        ("سلول موجودی خالی", f"برای {missing} کد، موجودی انبار پای کار در فایل خالی یا بدون ردیف بود؛ موقتاً صفر فرض شده و وضعیت هر ردیف مشخص است. خالی بودن لزوماً به معنای صفر واقعی نیست."),
        ("زمان موجودی", "زمان ایجاد فایل موجودی 2026-10-27 است؛ زمان دقیق ثبت خود موجودی در داده‌ها مشخص نیست. اگر مصرف برنامه قبلی پیش از ثبت موجودی انجام شده باشد، جمع کامل برنامه قبلی دوباره‌شماری ایجاد می‌کند."),
        ("درخواست باز / انبار مواد", "فایل‌های فعلی درخواست باز یا موجودی جاری انبار مواد اولیه را ندارند؛ ستون‌های مربوط عمداً خالی‌اند و از نمونه خروجی قدیمی به‌جای داده جاری استفاده نشده است."),
        ("زیرمجموعه‌های ۱۲", "قطعه‌های ۱۲ که در BOM محصول دیگری مصرف می‌شوند در خروجی آمده‌اند. تولید برنامه‌ریزی‌شده همان قطعه از درخواستش کم نشده، چون مقدار و زمان تکمیل واقعی معلوم نیست؛ قبل از درخواست انتقال یا خرید بررسی شود. BOM به‌صورت بازگشتی منفجر نشده تا برنامه مستقل زیرمجموعه دوباره‌شماری نشود."),
        ("سایر فرض‌ها", "ضایعات، درخواست‌های باز، بسته‌بندی، حداقل سفارش و تأمین از انبارهای دیگر لحاظ نشده‌اند. تغییر فایل‌های منبع خارج از این کتاب‌کار خودکار وارد نمی‌شود؛ برای بارگذاری داده‌های جدید، فایل باید دوباره تولید شود."),
        ("منابع", "bom(1).xlsx | برنامه 3روز.xlsx | برنامه بعدی.xlsx | موجودی به تفکیک انبار (1).xlsx؛ خروجی.xlsx فقط نمونه قالب است."),
    ]
    for index, (label, text) in enumerate(lines, 2):
        guide.cell(index, 1, label)
        guide.cell(index, 2, text)
        for cell in guide[index]:
            cell.font = Font(name="Tahoma", size=10, bold=(cell.column == 1), color="263445")
            cell.fill = PatternFill("solid", fgColor=PALE_YELLOW if label.startswith("هشدار") else ("F5F8FB" if index % 2 == 0 else WHITE))
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
            cell.border = BORDER
        guide.row_dimensions[index].height = 58 if len(text) > 180 else 42
    guide.freeze_panes = "B2"
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    return wb, cache


def main():
    data = build_data()
    wb, cache = create_workbook(data)
    save_with_cached_values(wb, OUTPUT_FILE, cache)
    positive = sum(request_amount(t["old"] + t["new"], data["machine"][c].quantity) > 0 for c, t in data["totals"].items())
    blank = sum(not data["machine"][c].saw_numeric for c in data["totals"])
    print(f"Saved: {OUTPUT_FILE}")
    print(f"Products: {len(data['products'])}; BOM lines: {len(data['detail'])}; components: {len(data['totals'])}; positive requests: {positive}; blank stock assumed zero: {blank}")


if __name__ == "__main__":
    main()
